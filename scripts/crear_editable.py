"""Copia editable del mismo contenido del informe, con texto y tablas nativos."""
import ast
import json
from html.parser import HTMLParser
from pathlib import Path
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import Paragraph, Spacer, Table, TableStyle, PageBreak, Image
from xml.sax.saxutils import escape

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'resultados'

def contenido():
    # Reutilizar las mismas instrucciones que componen el PDF, sin redibujar mapas.
    module=ast.parse((ROOT/'scripts'/'informe.py').read_text(encoding='utf-8'))
    fn=next(n for n in module.body if isinstance(n,ast.FunctionDef) and n.name=='generar')
    start=next(i for i,n in enumerate(fn.body) if isinstance(n,ast.Assign)
               and any(isinstance(t,ast.Name) and t.id=='styles' for t in n.targets))
    end=next(i for i,n in enumerate(fn.body) if isinstance(n,ast.FunctionDef) and n.name=='footer')
    segment=ast.Module(body=fn.body[start:end],type_ignores=[])
    s=json.loads((OUT/'resumen.json').read_text(encoding='utf-8'))
    context=dict(globals(),out=OUT,s=s,q=s['resultados'])
    exec(compile(segment,'contenido_original_pdf','exec'),context)
    return context['story']

class Rich(HTMLParser):
    def __init__(self,p):
        super().__init__();self.p=p;self.bold=False;self.italic=False
    def handle_starttag(self,tag,attrs):
        if tag in ('b','strong'):self.bold=True
        elif tag in ('i','em'):self.italic=True
        elif tag=='br':self.p.add_run().add_break()
    def handle_endtag(self,tag):
        if tag in ('b','strong'):self.bold=False
        elif tag in ('i','em'):self.italic=False
    def handle_data(self,data):
        r=self.p.add_run(data);r.bold=True if self.bold else None;r.italic=True if self.italic else None

def texto(p,html):Rich(p).feed(html)

def generar():
    doc=Document();sec=doc.sections[0]
    sec.page_width=Mm(210);sec.page_height=Mm(297)
    sec.top_margin=Pt(30);sec.bottom_margin=Pt(40)
    sec.left_margin=Pt(38);sec.right_margin=Pt(38)
    sec.footer_distance=Pt(18)
    normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(9)
    normal.font.color.rgb=RGBColor(0,0,0)
    pf=normal.paragraph_format;pf.line_spacing=Pt(12);pf.space_after=Pt(7)
    pf.widow_control=False
    for name,size in [('Title',21),('Heading 1',15),('Heading 2',13)]:
        st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size)
        st.font.color.rgb=RGBColor(0,0,0);st.font.bold=True
        for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
        for fonts in st.element.iter(qn('w:rFonts')):
            for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
                fonts.attrib.pop(qn('w:'+key),None)
        st.paragraph_format.line_spacing=Pt(size+3)
        st.paragraph_format.space_before=Pt(0 if name=='Title' else 8)
        st.paragraph_format.space_after=Pt(7)
    doc.styles['Title'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    blocks=contenido()
    for block in blocks:
        if isinstance(block,Paragraph):
            style={'Title':'Title','Heading1':'Heading 1','Heading2':'Heading 2'}.get(block.style.name,'Normal')
            p=doc.add_paragraph(style=style);texto(p,block.text)
        elif isinstance(block,PageBreak):doc.add_page_break()
        elif isinstance(block,Table):
            data=block._cellvalues; t=doc.add_table(rows=len(data),cols=len(data[0]))
            t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.LEFT
            widths=[float(w) for w in block._colWidths]
            ratio=519.27/sum(widths)
            for col,w in zip(t.columns,widths):col.width=Pt(w*ratio)
            borders=OxmlElement('w:tblBorders')
            for side in ('top','left','bottom','right','insideH','insideV'):
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single')
                el.set(qn('w:sz'),'3');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            t._tbl.tblPr.append(borders)
            for i,row in enumerate(data):
                trpr=t.rows[i]._tr.get_or_add_trPr()
                no_split=OxmlElement('w:cantSplit');trpr.append(no_split)
                if i==0:trpr.append(OxmlElement('w:tblHeader'))
                for j,value in enumerate(row):
                    cell=t.cell(i,j);cell.width=Pt(widths[j]*ratio)
                    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    pr=cell._tc.get_or_add_tcPr()
                    margins=OxmlElement('w:tcMar')
                    for side,size in [('top','32'),('bottom','32'),('left','80'),('right','80')]:
                        el=OxmlElement('w:'+side);el.set(qn('w:w'),size);el.set(qn('w:type'),'dxa');margins.append(el)
                    pr.append(margins)
                    if i==0:
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DBE8EF');pr.append(shade)
                    p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0)
                    p.paragraph_format.line_spacing=Pt(9.4)
                    p.paragraph_format.keep_with_next=False
                    texto(p,value.text if isinstance(value,Paragraph) else str(value))
                    for r in p.runs:r.font.size=Pt(7.8)
                    if i>0:
                        try:
                            float(p.text.replace(',',''))
                            p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
                        except ValueError:pass
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(2)
            p.add_run().font.size=Pt(2)
        elif isinstance(block,Image):
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(5)
            p.paragraph_format.line_spacing=1.0
            pic=p.add_run().add_picture(str(block.filename),width=Pt(519),height=Pt(239))
            pic._inline.docPr.set('descr',Path(block.filename).stem.replace('_',' '))
        # Los espaciadores de ReportLab ya se representan en los párrafos de Word.
    footer=sec.footer.paragraphs[0]
    footer.paragraph_format.space_after=Pt(0)
    r=footer.add_run('Galaz · Yañez · Roa | Santiago centro | 15-10-2026');r.font.size=Pt(8)
    footer.add_run(' '*9)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.core_properties.title='Ciudad construida y transporte público'
    doc.core_properties.author='Nicolas Galaz, Martin Yañez, Eduardo Roa'
    target=OUT/'Informe_Tarea1_PostGIS_editable.docx'
    doc.save(target)
    # Coincidencia literal del contenido de párrafos y celdas con el generador PDF.
    expected=[b.getPlainText() for b in blocks if isinstance(b,Paragraph)]
    actual=[p.text for p in doc.paragraphs if p.text]
    def norm(x):return ''.join(x.split())
    actualset={norm(x) for x in actual}
    assert all(norm(x) in actualset for x in expected)
    assert len(doc.tables)==sum(isinstance(b,Table) for b in blocks)
    assert len(doc.inline_shapes)==2
    originals=[b for b in blocks if isinstance(b,Table)]
    for editable,original in zip(doc.tables,originals):
        for row,source in zip(editable.rows,original._cellvalues):
            for cell,value in zip(row.cells,source):
                expected=value.getPlainText() if isinstance(value,Paragraph) else str(value)
                assert norm(cell.text)==norm(expected)
    print(target)

if __name__=='__main__':generar()
